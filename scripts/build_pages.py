"""키워드 원본 엑셀 → data/site-pages.json (사이트 라우트의 단일 출처) + keywords/선호환경_1020.xlsx

  python -I scripts/build_pages.py

원본: source/keywords_src.xlsx 첫 시트, A열 비어 있으면 무시. 헤더가 열과 어긋나 있어 실제 열 기준으로 읽는다.
  A 키워드 · B 붙여쓰기 · C 우선순위(1~3) · D 권역 · E 지역 표기 · F 지역 단위 · G 서비스 단어 · H 원문 · I 서비스 묶음 · J 수식어
키워드 1개 = 페이지 1개, 파일 순서 유지.
오픈 단계(wave) = 우선순위: 1 = 권역 대표 15 / 2 = 시·구·생활권 / 3 = 동·읍면
표기 변형(인천/인천시/인천광역시, 부평/부평구, 구월/구월동 …)은 같은 장소·같은 서비스의 대표(우선순위↑, 파일 앞) 페이지를 variantOf 로.
"""
import json, sys
from collections import Counter
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

sys.path.insert(0, str(Path(__file__).resolve().parent))
from places_src import place_key  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / 'source' / 'keywords_src.xlsx'
OUT_XLSX = ROOT / 'keywords' / '선호환경_1020.xlsx'
SITE_URL = 'https://sunho.example'  # TODO: 도메인 확정 후 lib/site.ts 와 함께 바꿀 것
UNIT = {'시·구': '시구', '생활권': '생활권', '동·읍면': '동'}

ws = load_workbook(SRC, read_only=True, data_only=True).worksheets[0]
rows = [r for r in list(ws.iter_rows(values_only=True))[1:] if r[0]]
pages = []
for i, r in enumerate(rows, 1):
    kw = str(r[0]).strip()
    pages.append(dict(n=i, kw=kw, slug=kw.replace(' ', '-'), label=r[4], place=place_key(r[4]), region=r[3], unit=UNIT[r[5]],
                      service=r[6], group=r[8], mod=r[9] or '', prio=int(r[2]), wave=int(r[2])))
assert len({p['slug'] for p in pages}) == len(pages), '슬러그 중복'

# 표기 변형 → 대표
best = {}
for p in pages:
    k = (p['place'], p['service'])
    if k not in best or (p['prio'], p['n']) < (best[k]['prio'], best[k]['n']):
        best[k] = p
for p in pages:
    b = best[(p['place'], p['service'])]
    if b is not p:
        p['variantOf'] = b['slug']

data = {'pages': pages}
(ROOT / 'data' / 'site-pages.json').write_text(json.dumps(data, ensure_ascii=False, indent=0), encoding='utf-8')

# 작업용 엑셀 (키워드 시트가 첫 시트)
wb = Workbook(); sh = wb.active; sh.title = '키워드'
head = ['번호', '키워드', '우선순위(단계)', '권역', '지역 표기', '장소', '지역 단위', '서비스', '수식어', 'URL', '대표 페이지(표기 변형)']
sh.append(head)
for p in pages:
    sh.append([p['n'], p['kw'], p['wave'], p['region'], p['label'], p['place'], p['unit'], p['service'], p['mod'],
               f"{SITE_URL}/{p['slug']}/", p.get('variantOf', '')])
for c in sh[1]:
    c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='102A2A'); c.alignment = Alignment(horizontal='center')
for col, w in zip('ABCDEFGHIJK', [6, 28, 8, 8, 14, 12, 8, 10, 6, 50, 28]):
    sh.column_dimensions[col].width = w
sh.freeze_panes = 'A2'
OUT_XLSX.parent.mkdir(exist_ok=True)
wb.save(OUT_XLSX)

reps = [p for p in pages if 'variantOf' not in p]
print('페이지', len(pages), '대표', len(reps), '변형', len(pages) - len(reps))
for w in (1, 2, 3):
    op = [p for p in pages if p['wave'] <= w]
    print(f'{w}단계: 열린 {len(op)} / 사이트맵 키워드 {sum(1 for p in op if "variantOf" not in p or best[(p["place"], p["service"])]["wave"] > w)}')
print(Counter(p['service'] for p in pages))
